import os
import re
import sys
from datetime import datetime

import pdfplumber
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# === CONFIGURATION ===
DEFAULT_INPUT_FOLDER = r"C:\Users\anish\Downloads\PanSo_Data"
OUTPUT_FOLDER = r"C:\Users\anish\Downloads"


def extract_first_match(pattern, text, flags=re.MULTILINE):
    m = re.search(pattern, text, flags)
    return m.group(1).strip() if m else ""


def extract_header_fields(text):
    fiscal_year = extract_first_match(r"Fiscal Year:\s*(\d{4})", text)
    payment_month = extract_first_match(r"([A-Za-z]+)\s+Payment", text)

    m = re.search(r"Name:\s*(.*?)\s+County:\s*(.*?)\s+IRN:\s*([0-9]+)", text)
    if m:
        school_name, county, irn = m.group(1).strip(), m.group(2).strip(), m.group(3).strip()
    else:
        school_name = county = irn = ""

    sponsor_name = extract_first_match(r"Sponsor Name:\s*(.*?)\s+E-School:", text)

    return {
        "School Name": school_name,
        "County": county,
        "IRN": irn,
        "Sponsor Name": sponsor_name,
        "Fiscal Year": fiscal_year,
        "Payment Month": payment_month,
    }


def extract_last_amount_from_line(line):
    numbers = re.findall(r"-?\d[\d,]*\.\d{2}", line)
    return numbers[-1] if numbers else ""


def extract_monthly_value(label, text):
    for line in text.splitlines():
        if line.strip().startswith(label):
            return extract_last_amount_from_line(line)
    return ""


def extract_financial_fields(text):
    label_map = {
        "Base Cost Monthly": "Base Cost",
        "Special Education Monthly": "Special Education",
        "DPIA Monthly": "Disadvantaged Pupil Impact Aid (DPIA)",
        "Facilities Funding Monthly": "Facilities Funding",
        "Total State Support Monthly": "Total State Support",
        "Total Transfers Monthly": "Total Transfers",
        "Total Adjustments Monthly": "Total Adjustments",
        "Total Payment Monthly": "Total Payment",
    }
    result = {}
    for field_name, pdf_label in label_map.items():
        result[field_name] = extract_monthly_value(pdf_label, text)
    return result


def process_pdf(pdf_path):
    print(f"Processing: {pdf_path}")
    try:
        with pdfplumber.open(pdf_path) as pdf:
            full_text = "\n".join((page.extract_text() or "") for page in pdf.pages)

        header = extract_header_fields(full_text)
        financials = extract_financial_fields(full_text)

        combined = {"PDF Filename": os.path.basename(pdf_path)}
        combined.update(header)
        combined.update(financials)

        return combined

    except Exception as e:
        print(f"ERROR reading {pdf_path}: {e}")
        return None


def format_excel(path):
    wb = load_workbook(path)
    ws = wb.active

    bold = Font(bold=True)
    header_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    section_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    thin = Side(border_style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    max_row = ws.max_row

    current_school = ""

    for row in range(1, max_row + 1):

        field_cell = ws[f"A{row}"]
        value_cell = ws[f"B{row}"]

        # Add borders and alignment
        field_cell.border = border
        value_cell.border = border
        field_cell.alignment = Alignment(vertical="center")
        value_cell.alignment = Alignment(vertical="center")

        # Make Field column bold
        field_cell.font = bold

        # Detect a new school section
        if field_cell.value == "School Name":
            current_school = value_cell.value  # School name appears in "Value" column

            # Style section header
            section_header_cell = ws[f"A{row}"]
            section_value_cell = ws[f"B{row}"]

            section_header_cell.fill = section_fill
            section_value_cell.fill = section_fill
            section_header_cell.font = Font(bold=True)
            section_value_cell.font = Font(bold=True)

        # PDF Filename row formatting
        if field_cell.value == "PDF Filename":
            field_cell.fill = header_fill
            value_cell.fill = header_fill

    wb.save(path)


def main(input_folder):
    if not os.path.isdir(input_folder):
        print(f"Input folder not found: {input_folder}")
        sys.exit(1)

    all_records = []

    for file in os.listdir(input_folder):
        if file.lower().endswith(".pdf"):
            full_path = os.path.join(input_folder, file)
            record = process_pdf(full_path)
            if record:
                all_records.append(record)

    if not all_records:
        print("No PDF data extracted.")
        return

    # Build vertical layout for Excel
    rows = []
    for record in all_records:
        # Add PDF filename first
        rows.append({"Field": "PDF Filename", "Value": record["PDF Filename"]})

        for k, v in record.items():
            if k != "PDF Filename":
                rows.append({"Field": k, "Value": v})

        rows.append({"Field": "", "Value": ""})  # separator

    df = pd.DataFrame(rows)

    timestamp = datetime.now().strftime("%Y-%m-%d_%I-%M-%p")
    output_path = os.path.join(OUTPUT_FOLDER, f"PanSo_Settlements_{timestamp}.xlsx")

    df.to_excel(output_path, index=False)
    format_excel(output_path)

    print(f"\n✅ Done! Excel created at:\n{output_path}\n")


if __name__ == "__main__":
    folder = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_INPUT_FOLDER
    main(folder)
